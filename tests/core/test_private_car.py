"""私车公用 core 模块测试：口径计算、行校验、模板导出。"""

from datetime import datetime, time
from pathlib import Path
from unittest.mock import patch

import pytest
from invoice_processor.core.private_car import (
    SUBSIDY_PER_KM,
    calculate_private_car_cost,
    export_private_car_excel,
    normalize_row,
    validate_row,
)
from openpyxl import load_workbook

TEMPLATE = Path(__file__).resolve().parents[2] / (
    'backend/invoice_processor/core/templates/private_car_v1.xlsx'
)


def _row(**overrides):
    row = dict(
        date='2026-10-23', user='李伟', destination='客户甲',
        start_time='9:00', end_time='17:40',
        start_mileage='44118', end_mileage='44165',
        toll='88', parking='10',
    )
    row.update(overrides)
    return row


class TestCalculatePrivateCarCost:
    def test_subsidy_per_km_is_one_point_five(self):
        assert SUBSIDY_PER_KM == 1.5

    def test_distance_and_subsidy(self):
        distance, subsidy, error = calculate_private_car_cost('44118', '44165')
        assert (distance, subsidy, error) == ('47', '70.50', '')

    def test_rounded_subsidy_two_decimals(self):
        distance, subsidy, error = calculate_private_car_cost('0', '10.5')
        assert (distance, subsidy, error) == ('10.5', '15.75', '')

    def test_end_not_less_than_start(self):
        _, _, error = calculate_private_car_cost('100', '90')
        assert error == '结束里程数不能小于出发里程数'

    def test_invalid_number(self):
        _, _, error = calculate_private_car_cost('abc', '200')
        assert '里程数' in error


class TestValidateRow:
    def test_complete_row_passes(self):
        assert validate_row(_row()) == ''

    def test_missing_user(self):
        assert '使用人' in validate_row(_row(user=' '))

    def test_bad_date(self):
        assert validate_row(_row(date='10/23')) != ''

    def test_bad_time(self):
        assert validate_row(_row(start_time='9点')) != ''

    def test_end_mileage_before_start(self):
        assert validate_row(_row(start_mileage='500', end_mileage='100')) != ''

    def test_fee_not_negative(self):
        assert validate_row(_row(toll='-1')) != ''

    def test_optional_fees_default_zero(self):
        assert validate_row(_row(toll='', parking='')) == ''


class TestNormalizeRow:
    def test_pads_time_and_keeps_fees_as_is(self):
        normalized = normalize_row(_row(start_time='9:00', toll='', parking=''))
        assert normalized['start_time'] == '09:00'
        assert normalized['toll'] == ''
        assert normalized['parking'] == ''


class TestExportPrivateCarExcel:
    def export(self, tmp_path, rows, **kwargs):
        return export_private_car_excel(
            rows, license_plate='苏U930FQ',
            source_directory=str(tmp_path / '客户甲'),
            output_directory=str(tmp_path),
            template_path=TEMPLATE, **kwargs,
        )

    def test_layout_values_and_logo_preserved(self, tmp_path):
        original = TEMPLATE.read_bytes()
        rows = [normalize_row(_row())]
        path = self.export(tmp_path, rows)
        assert path.name == '新代私车公用费用统计表_2026-10-23_客户甲.xlsx'
        workbook = load_workbook(path)
        try:
            sheet = workbook.active
            assert sheet['K2'].value == '苏U930FQ'
            assert sheet['B4'].value == datetime(2026, 10, 23)
            assert sheet['B4'].number_format == 'yyyy-mm-dd'
            assert sheet['C4'].value == '李伟'
            assert sheet['D4'].value == '客户甲'
            assert sheet['E4'].value == time(9, 0)
            assert sheet['F4'].value == time(17, 40)
            assert sheet['E4'].number_format == 'h:mm'
            assert [sheet.cell(4, col).value for col in range(7, 11)] == \
                [44118, 44165, 47, 70.5]
            assert sheet['K4'].value == '过路（桥）费：88'
            assert sheet['K5'].value == '停车费：10'
            assert 'B4:B5' in {str(value) for value in sheet.merged_cells.ranges}
            assert sheet['B26'].value == '*CR-GF01-00-06 v1'
            assert len(sheet._images) == 1
            assert sheet.page_setup.orientation == 'landscape'
        finally:
            workbook.close()
        # 模板原件永不修改
        assert TEMPLATE.read_bytes() == original

    def test_more_than_eight_records_use_additional_sheets(self, tmp_path):
        rows = [normalize_row(_row()) for _ in range(9)]
        rows[-1] = normalize_row(_row(
            date='2026-10-24', start_mileage='10', end_mileage='20',
            toll='1', parking='2'))
        path = self.export(tmp_path, rows)
        assert '2026-10-23至2026-10-24' in path.name
        workbook = load_workbook(path)
        try:
            assert len(workbook.worksheets) == 2
            first, second = workbook.worksheets
            assert first['B18'].value == datetime(2026, 10, 23)
            assert second['B4'].value == datetime(2026, 10, 24)
            assert second['J4'].value == 15
            assert second['B6'].value is None
            assert all(len(sheet._images) == 1 for sheet in workbook)
            assert second['K2'].value == '苏U930FQ'
        finally:
            workbook.close()

    def test_formula_injection_is_stored_as_text(self, tmp_path):
        rows = [normalize_row(_row(destination='=1+1'))]
        path = self.export(tmp_path, rows)
        workbook = load_workbook(path)
        try:
            assert workbook.active['D4'].data_type == 's'
        finally:
            workbook.close()

    def test_failed_replace_keeps_old_file(self, tmp_path):
        rows = [normalize_row(_row())]
        path = self.export(tmp_path, rows)
        original = path.read_bytes()
        with patch('invoice_processor.core.private_car.os.replace',
                   side_effect=PermissionError('Excel占用')):
            with pytest.raises(PermissionError):
                self.export(tmp_path, rows)
        assert path.read_bytes() == original
        assert list(tmp_path.glob('*.xlsx')) == [path]

    def test_empty_rows_and_missing_template_are_reported(self, tmp_path):
        with pytest.raises(ValueError):
            self.export(tmp_path, [])
        with pytest.raises(FileNotFoundError):
            export_private_car_excel(
                [normalize_row(_row())], license_plate='苏U930FQ',
                source_directory=str(tmp_path / '客户甲'),
                output_directory=str(tmp_path),
                template_path=tmp_path / 'missing.xlsx',
            )


class TestTemplateAsset:
    def test_builtin_template_ships_with_package(self):
        assert TEMPLATE.is_file()
        assert TEMPLATE.name == 'private_car_v1.xlsx'

    def test_template_copies_once_into_app_directory(self, tmp_path, monkeypatch):
        import invoice_processor.core.private_car as pc

        monkeypatch.setattr(pc, 'TEMPLATE_ASSET', TEMPLATE)
        target_dir = tmp_path / 'app'
        target_dir.mkdir()
        monkeypatch.setattr(pc, 'get_app_directory', lambda: target_dir)

        first = pc.get_private_car_template()
        assert first == target_dir / pc.TEMPLATE_FILE_NAME
        assert first.read_bytes() == TEMPLATE.read_bytes()
        # 第二次调用不再覆盖副本
        first.write_bytes(b'customized')
        second = pc.get_private_car_template()
        assert second.read_bytes() == b'customized'
