"""单元测试：AI 审核模块（提示词构造 + 响应解析 + 报告写入防注入，不联网）

运行方式: pytest tests/test_ai_audit.py -v
"""
import json
import os
import urllib.request

from invoice_processor.core.ai_audit import (
    _sanitize_cell,
    build_prompt,
    parse_findings,
    write_audit_report,
)
from invoice_processor.core.ai_audit import test_connection as check_ai_connection
from openpyxl import Workbook, load_workbook


class FakeResponse:
    def __init__(self, body):
        self.body = body

    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc_value, traceback):
        return False

    def read(self):
        return self.body


class TestBuildPrompt:
    def test_contains_data(self):
        records = [{
            'file': 'a.pdf',
            'rows': [{'date': '2026-07-27', 'category': 'transport',
                      'transport_amount': 17.9}],
        }]
        prompt = build_prompt(records)
        assert 'a.pdf' in prompt
        assert '2026-07-27' in prompt

    def test_chinese_audit_rules(self):
        prompt = build_prompt([{'file': 'x'}])
        assert '金额' in prompt
        assert '重复' in prompt
        # 审核重点是金额错误，行程只做简易核对
        assert '重点是金额错误' in prompt


class TestParseFindings:
    def test_empty_array(self):
        assert parse_findings('[]') == []

    def test_plain_json_array(self):
        content = (
            '[{"file": "a.pdf", "type": "conflict", '
            '"issue": "时间冲突", "suggestion": "核对"}]'
        )
        out = parse_findings(content)
        assert len(out) == 1
        assert out[0]['file'] == 'a.pdf'
        assert out[0]['type'] == 'conflict'

    def test_markdown_fenced(self):
        content = (
            '```json\n[{"file": "b.pdf", "type": "extract", '
            '"issue": "x", "suggestion": "y"}]\n```'
        )
        out = parse_findings(content)
        assert len(out) == 1
        assert out[0]['type'] == 'extract'

    def test_noise_around(self):
        content = (
            '好的，以下是审核结果：\n'
            '[{"file": "c.pdf", "type": "duplicate", "issue": "重复", '
            '"suggestion": "z"}]\n完毕'
        )
        out = parse_findings(content)
        assert len(out) == 1
        assert out[0]['type'] == 'duplicate'

    def test_invalid_inputs(self):
        assert parse_findings('无法解析') == []
        assert parse_findings(None) == []
        assert parse_findings('') == []
        assert parse_findings('{}') == []


def test_connection_posts_minimal_chat_request(monkeypatch):
    captured = {}

    def fake_urlopen(request, timeout):
        captured['request'] = request
        captured['timeout'] = timeout
        return FakeResponse(b'{"choices":[{"message":{"content":"OK"}}]}')

    monkeypatch.setattr(urllib.request, 'urlopen', fake_urlopen)

    check_ai_connection(
        'pending-key',
        api_base='https://ai.example.com/',
        model='test-model',
        timeout=15,
    )

    request = captured['request']
    assert request.full_url == 'https://ai.example.com/chat/completions'
    assert request.get_header('Authorization') == 'Bearer pending-key'
    assert captured['timeout'] == 15
    assert json.loads(request.data) == {
        'model': 'test-model',
        'messages': [{'role': 'user', 'content': '请仅回复 OK'}],
        'temperature': 0,
        'max_tokens': 1,
        'stream': False,
    }


class TestSanitizeCell:
    def test_formula_prefix_escaped(self):
        for prefix in ('=', '+', '-', '@', '\t', '\r'):
            assert _sanitize_cell(f'{prefix}HYPERLINK("x")') == (
                f"'{prefix}HYPERLINK(\"x\")"
            )

    def test_plain_values_untouched(self):
        assert _sanitize_cell('打车单程 150.00 元超标') == '打车单程 150.00 元超标'
        assert _sanitize_cell('AI 审核') == 'AI 审核'
        assert _sanitize_cell(None) == ''
        assert _sanitize_cell(123) == '123'


class TestWriteAuditReport:
    def _xlsx_with_summary(self, tmp_path):
        """费用汇总.xlsx 存在是 write_audit_report 的前置条件"""
        wb = Workbook()
        ws = wb.active
        ws.title = '费用汇总'
        path = os.path.join(str(tmp_path), '费用汇总.xlsx')
        wb.save(path)
        return path

    def test_formula_finding_neutralized(self, tmp_path):
        """AI findings 引用 PDF 文本（=HYPERLINK 开头）→ 写入时加 ' 前缀"""
        xlsx = self._xlsx_with_summary(tmp_path)
        findings = [{
            'source': 'AI 审核',
            'file': '=HYPERLINK("http://evil", "点此")',
            'type': 'other',
            'issue': '=cmd|\' /C calc\'!A0',
            'suggestion': '建议',
        }]
        result = write_audit_report(str(tmp_path), findings)

        assert result == xlsx
        wb = load_workbook(xlsx)
        ws = wb['审核报告']
        row = [ws.cell(row=2, column=c).value for c in range(1, 6)]
        assert row[1] == "'=HYPERLINK(\"http://evil\", \"点此\")"
        assert row[3] == "'=cmd|' /C calc'!A0"
        # openpyxl data_type 不是 f（formula）
        assert ws.cell(row=2, column=4).data_type != 'f'

    def test_normal_finding_written_as_is(self, tmp_path):
        xlsx = self._xlsx_with_summary(tmp_path)
        findings = [{
            'source': '本地规则',
            'file': 'a.pdf',
            'type': 'other',
            'issue': '打车单程 150.00 元超标',
            'suggestion': '附超标说明',
        }]
        write_audit_report(str(tmp_path), findings)
        wb = load_workbook(xlsx)
        ws = wb['审核报告']
        assert ws.cell(row=2, column=4).value == '打车单程 150.00 元超标'

    def test_no_summary_returns_none(self, tmp_path):
        assert write_audit_report(str(tmp_path), []) is None

    def test_empty_findings_writes_ok_row(self, tmp_path):
        xlsx = self._xlsx_with_summary(tmp_path)
        write_audit_report(str(tmp_path), [])
        wb = load_workbook(xlsx)
        ws = wb['审核报告']
        assert ws.cell(row=2, column=4).value == '无异常'
